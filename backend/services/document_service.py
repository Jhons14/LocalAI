"""
Document processing service for handling uploaded files.
"""
import io
import tempfile
from typing import Dict, Any, Optional, Tuple
from pathlib import Path
import logging
import mimetypes

# PDF processing
try:
    from pypdf import PdfReader
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

# DOCX processing
try:
    from docx import Document as DocxDocument
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Service for processing uploaded documents and extracting text content."""
    
    ALLOWED_EXTENSIONS = {'.txt', '.md', '.pdf', '.docx'}
    ALLOWED_MIME_TYPES = {
        'text/plain',
        'text/markdown', 
        'application/pdf',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    }
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    MAX_CONTENT_LENGTH = 50000  # Maximum characters to include in prompt
    
    @classmethod
    def validate_file(cls, filename: str, content: bytes) -> Dict[str, Any]:
        """Validate file type, size, and basic security checks."""
        validation_result = {
            'valid': False,
            'error': None,
            'file_info': {}
        }
        
        try:
            # Check file size
            file_size = len(content)
            if file_size > cls.MAX_FILE_SIZE:
                validation_result['error'] = f"File size ({file_size / 1024 / 1024:.1f}MB) exceeds maximum ({cls.MAX_FILE_SIZE / 1024 / 1024}MB)"
                return validation_result
            
            # Check file extension
            file_extension = Path(filename).suffix.lower()
            if file_extension not in cls.ALLOWED_EXTENSIONS:
                validation_result['error'] = f"File extension '{file_extension}' not allowed. Supported: {', '.join(cls.ALLOWED_EXTENSIONS)}"
                return validation_result
            
            # Check MIME type (basic check)
            mime_type, _ = mimetypes.guess_type(filename)
            if mime_type and mime_type not in cls.ALLOWED_MIME_TYPES:
                validation_result['error'] = f"MIME type '{mime_type}' not allowed"
                return validation_result
            
            # Basic security check - look for suspicious content in first 1000 bytes
            sample = content[:1000].decode('utf-8', errors='ignore').lower()
            suspicious_patterns = ['<script', '<?php', 'eval(', 'exec(', 'system(']
            for pattern in suspicious_patterns:
                if pattern in sample:
                    validation_result['error'] = f"Potentially malicious content detected"
                    return validation_result
            
            validation_result['valid'] = True
            validation_result['file_info'] = {
                'filename': filename,
                'size': file_size,
                'extension': file_extension,
                'mime_type': mime_type
            }
            
        except Exception as e:
            validation_result['error'] = f"File validation error: {str(e)}"
            
        return validation_result
    
    @classmethod
    def extract_text_content(cls, filename: str, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """Extract text content from various file types."""
        try:
            file_extension = Path(filename).suffix.lower()
            extracted_text = ""
            metadata = {
                'extraction_method': 'unknown',
                'char_count': 0,
                'truncated': False
            }
            
            if file_extension in ['.txt', '.md']:
                # Handle text files
                try:
                    extracted_text = content.decode('utf-8')
                    metadata['extraction_method'] = 'utf-8_decode'
                except UnicodeDecodeError:
                    try:
                        extracted_text = content.decode('latin-1')
                        metadata['extraction_method'] = 'latin-1_decode'
                    except UnicodeDecodeError:
                        extracted_text = content.decode('utf-8', errors='ignore')
                        metadata['extraction_method'] = 'utf-8_decode_ignore_errors'
                        
            elif file_extension == '.pdf':
                extracted_text, metadata = cls._extract_pdf_content(content, metadata)

            elif file_extension == '.docx':
                extracted_text, metadata = cls._extract_docx_content(content, metadata)
                
            else:
                raise ValueError(f"Unsupported file type: {file_extension}")
            
            # Truncate content if too long
            original_length = len(extracted_text)
            if original_length > cls.MAX_CONTENT_LENGTH:
                extracted_text = extracted_text[:cls.MAX_CONTENT_LENGTH]
                metadata['truncated'] = True
                metadata['original_length'] = original_length
            
            metadata['char_count'] = len(extracted_text)
            
            return extracted_text, metadata
            
        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {str(e)}")
            raise ValueError(f"Failed to extract text content: {str(e)}")

    @classmethod
    def _extract_pdf_content(cls, content: bytes, metadata: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
        """Extract text content from PDF files using pypdf."""
        if not PDF_SUPPORT:
            logger.warning("pypdf not installed - PDF extraction unavailable")
            return "[PDF processing requires pypdf library - please install it or use text files]", metadata

        try:
            pdf_file = io.BytesIO(content)
            reader = PdfReader(pdf_file)

            text_parts = []
            page_count = len(reader.pages)
            metadata['page_count'] = page_count

            for page_num, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"[Page {page_num + 1}]\n{page_text}")
                except Exception as e:
                    logger.warning(f"Failed to extract text from page {page_num + 1}: {e}")
                    text_parts.append(f"[Page {page_num + 1}] - Could not extract text")

            extracted_text = "\n\n".join(text_parts)
            metadata['extraction_method'] = 'pypdf'

            if not extracted_text.strip():
                # PDF might be scanned/image-based
                extracted_text = "[PDF appears to be image-based or contains no extractable text]"
                metadata['extraction_method'] = 'pypdf_no_text'

            return extracted_text, metadata

        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            raise ValueError(f"Failed to extract PDF content: {str(e)}")

    @classmethod
    def _extract_docx_content(cls, content: bytes, metadata: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
        """Extract text content from DOCX files using python-docx."""
        if not DOCX_SUPPORT:
            logger.warning("python-docx not installed - DOCX extraction unavailable")
            return "[DOCX processing requires python-docx library - please install it or use text files]", metadata

        try:
            docx_file = io.BytesIO(content)
            doc = DocxDocument(docx_file)

            text_parts = []

            # Extract text from paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)

            # Extract text from tables
            table_count = len(doc.tables)
            if table_count > 0:
                metadata['table_count'] = table_count
                for table_idx, table in enumerate(doc.tables):
                    table_text = []
                    for row in table.rows:
                        row_text = [cell.text.strip() for cell in row.cells]
                        table_text.append(" | ".join(row_text))
                    if table_text:
                        text_parts.append(f"\n[Table {table_idx + 1}]\n" + "\n".join(table_text))

            extracted_text = "\n\n".join(text_parts)
            metadata['extraction_method'] = 'python-docx'
            metadata['paragraph_count'] = len(doc.paragraphs)

            if not extracted_text.strip():
                extracted_text = "[DOCX document appears to be empty or contains no extractable text]"
                metadata['extraction_method'] = 'python-docx_no_text'

            return extracted_text, metadata

        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            raise ValueError(f"Failed to extract DOCX content: {str(e)}")

    @classmethod
    def process_document(cls, filename: str, content: bytes) -> Dict[str, Any]:
        """Complete document processing pipeline."""
        result = {
            'success': False,
            'text_content': '',
            'metadata': {},
            'error': None
        }
        
        try:
            # Validate file
            validation = cls.validate_file(filename, content)
            if not validation['valid']:
                result['error'] = validation['error']
                return result
            
            # Extract text content
            text_content, extraction_metadata = cls.extract_text_content(filename, content)
            
            result.update({
                'success': True,
                'text_content': text_content,
                'metadata': {
                    'file_info': validation['file_info'],
                    'extraction': extraction_metadata,
                    'processing_timestamp': None  # Could add timestamp if needed
                }
            })
            
        except Exception as e:
            result['error'] = str(e)
            logger.error(f"Document processing failed for {filename}: {str(e)}")
            
        return result
    
    @classmethod
    def format_document_for_llm(cls, user_message: str, document_content: str, filename: str) -> str:
        """Format the user message and document content for optimal LLM processing."""
        if not document_content.strip():
            return user_message
            
        formatted_prompt = f"""{user_message}

--- Document Content (File: {filename}) ---
{document_content.strip()}
--- End of Document ---

Please analyze the above document content in the context of my request."""
        
        return formatted_prompt
    
    @classmethod
    def get_supported_formats(cls) -> Dict[str, str]:
        """Get list of supported file formats with descriptions."""
        pdf_status = "full support" if PDF_SUPPORT else "requires pypdf"
        docx_status = "full support" if DOCX_SUPPORT else "requires python-docx"
        return {
            '.txt': 'Plain text files',
            '.md': 'Markdown files',
            '.pdf': f'PDF documents ({pdf_status})',
            '.docx': f'Microsoft Word documents ({docx_status})'
        }